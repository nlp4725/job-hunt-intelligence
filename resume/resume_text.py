"""
Uploaded resume file -> normalized text for extract_skills().

The resume-side adapter in front of the shared extractor: it only decides how
to read each file type; both resume and JD text then go through the same
normalize() and the same taxonomy. See docs/resume_jd_skill_pipeline.md §4.1.
"""

import io
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

import docx
import pdfplumber
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from pdfplumber.utils.exceptions import PdfminerException

from analysis.text_normalize import normalize


class UnsupportedResume(Exception):
    """The file can't be turned into text; the message is shown to the user."""


@dataclass
class ResumeDocument:
    text: str


def resume_to_text(filename: str, data: bytes) -> ResumeDocument:
    suffix = PurePath(filename).suffix.lower()
    if suffix in (".txt", ".md"):
        return ResumeDocument(text=normalize(data.decode("utf-8")))
    if suffix == ".docx":
        return ResumeDocument(text=normalize(_docx_text(data)))
    if suffix == ".pdf":
        return ResumeDocument(text=normalize(_pdf_text(data)))
    raise UnsupportedResume(f"{suffix or 'This file type'} isn't supported. Upload a PDF, DOCX, TXT or MD file.")


def _column_gutter(page, min_gap: float = 15) -> tuple[int, int] | None:
    """The widest vertical strip in the middle half of the page that no word
    crosses on any line, or None for a single-column page.

    Whole-page on purpose: a single-column resume always has some full-width
    line (a bullet, the summary) crossing the middle, even when a role and a
    right-aligned date share a row. A false split can only happen where no
    word crosses, so it reorders text but never cuts a word or a phrase."""
    width = int(page.width)
    covered = [False] * (width + 1)
    for word in page.extract_words():
        for x in range(int(word["x0"]), min(width, int(word["x1"])) + 1):
            covered[x] = True
    lo, hi = int(width * 0.25), int(width * 0.75)
    best, start = None, None
    for x in range(lo, hi + 1):
        if x < hi and not covered[x]:
            start = x if start is None else start
            continue
        if start is not None and x - start >= min_gap and (best is None or x - start > best[1] - best[0]):
            best = (start, x)
        start = None
    return best


def _page_text(page) -> str:
    """pdfplumber reads straight across the page, which glues a left-column
    line to the right column's line on the same row. Read each column in turn."""
    gutter = _column_gutter(page)
    if gutter is None:
        return page.extract_text() or ""
    middle = sum(gutter) / 2
    left = page.crop((0, 0, middle, page.height)).extract_text() or ""
    right = page.crop((middle, 0, page.width, page.height)).extract_text() or ""
    return left + "\n" + right


def _pdf_text(data: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join(_page_text(page) for page in pdf.pages)
    except PdfminerException as e:
        raise UnsupportedResume("This PDF couldn't be opened. Try exporting it again.") from e
    if not text.strip():
        # A scan is an image with no text layer. No OCR by design: refusing is
        # better than silently scoring the resume as having no skills.
        raise UnsupportedResume("This PDF has no selectable text (it looks like a scan). Upload a PDF exported from your editor, or a DOCX.")
    return text


def _docx_text(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError) as e:
        raise UnsupportedResume("This DOCX file couldn't be opened. Try saving it again or upload a PDF.") from e

    # Body order, tables included: many templates put the skills section in a
    # table, and document.paragraphs skips table cells entirely.
    lines = []
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                # " | " rather than a space, so words in neighbouring cells
                # can't form a multi-word skill that isn't there.
                lines.append(" | ".join(cell.text for cell in row.cells))
        else:
            lines.append(block.text)
    return "\n".join(lines)
