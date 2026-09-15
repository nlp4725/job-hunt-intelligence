"""resume_to_text() turns an uploaded resume file into normalized text that
extract_skills() reads. Every test asserts on what the pipeline needs from
it: the skills survive the conversion, and files it can't read are refused
with UnsupportedResume rather than silently producing empty text.
"""

import io

import docx
import pytest
from fpdf import FPDF

from analysis.skills_extractor import extract_skills
from resume.resume_text import UnsupportedResume, resume_to_text

RESUME = (
    "SKILLS\n"
    "Python, SQL, PyTorch, Docker, Prompt\n"
    "Engineering, fine‑tuning\n"
    "EXPERIENCE\n"
    "Built RAG systems with LangGraph on AWS.\n"
)
EXPECTED = {"Python", "SQL", "PyTorch", "Docker", "Prompt Engineering",
            "Fine-tuning", "RAG", "LangGraph", "AWS"}


class TestPlainText:
    @pytest.mark.parametrize("filename", ["resume.txt", "Resume.MD"])
    def test_text_resume_keeps_its_skills(self, filename):
        doc = resume_to_text(filename, RESUME.encode("utf-8"))
        assert EXPECTED <= set(extract_skills(doc.text))

    def test_text_is_normalized(self):
        """The same normalize() that JDs go through, so a non-breaking
        hyphen in the file doesn't reach the extractor."""
        doc = resume_to_text("resume.txt", RESUME.encode("utf-8"))
        assert "‑" not in doc.text

    def test_unknown_file_type_is_refused(self):
        with pytest.raises(UnsupportedResume):
            resume_to_text("resume.pages", b"anything")


def _docx_bytes(text: str) -> bytes:
    """Render resume text as a DOCX: upper-case lines become Heading 1,
    everything else a normal paragraph, the way a resume template looks."""
    document = docx.Document()
    for line in text.splitlines():
        if line.isupper():
            document.add_heading(line, level=1)
        else:
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestDocx:
    def test_docx_resume_keeps_the_same_skills_as_its_text(self):
        doc = resume_to_text("resume.docx", _docx_bytes(RESUME))
        assert EXPECTED <= set(extract_skills(doc.text))

    def test_skills_inside_a_table_are_found(self):
        """Many resume templates lay the skills section out as a table, and
        python-docx's document.paragraphs skips table cells entirely."""
        document = docx.Document()
        document.add_heading("SKILLS", level=1)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Languages"
        table.cell(0, 1).text = "Python, SQL"
        table.cell(1, 0).text = "ML"
        table.cell(1, 1).text = "PyTorch, LangGraph"
        document.add_heading("EXPERIENCE", level=1)
        document.add_paragraph("Deployed on AWS.")
        buffer = io.BytesIO()
        document.save(buffer)

        found = set(extract_skills(resume_to_text("resume.docx", buffer.getvalue()).text))
        assert {"Python", "SQL", "PyTorch", "LangGraph", "AWS"} <= found

    def test_corrupt_docx_is_refused(self):
        with pytest.raises(UnsupportedResume):
            resume_to_text("resume.docx", b"not really a docx")


def _pdf_bytes(text: str, width_mm: float = 60) -> bytes:
    """Render resume text as a PDF. The narrow text width makes lines wrap in
    the middle of skill phrases, as real resume PDFs do ("Vector\\nDatabase").
    fpdf2's built-in fonts are Latin-1 only, hence the hyphen swap."""
    pdf = FPDF()
    pdf.add_page()
    for line in text.replace("‑", "-").splitlines():
        pdf.set_font("Helvetica", "B" if line.isupper() else "", 14 if line.isupper() else 11)
        pdf.multi_cell(width_mm, 6, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


class TestPdf:
    def test_pdf_resume_keeps_the_same_skills_as_its_text(self):
        doc = resume_to_text("resume.pdf", _pdf_bytes(RESUME))
        assert EXPECTED <= set(extract_skills(doc.text))

    def test_pdf_with_no_text_layer_is_refused(self):
        """A scanned resume is an image: there is no text to extract, and no
        OCR by design. Refusing beats silently scoring it as zero skills."""
        pdf = FPDF()
        pdf.add_page()
        pdf.rect(20, 20, 150, 200, style="F")
        with pytest.raises(UnsupportedResume, match="no selectable text"):
            resume_to_text("scan.pdf", bytes(pdf.output()))

    def test_corrupt_pdf_is_refused(self):
        with pytest.raises(UnsupportedResume, match="couldn't be opened"):
            resume_to_text("resume.pdf", b"not really a pdf")

    def test_two_column_pdf_is_read_column_by_column(self):
        """pdfplumber reads straight across the page by default, gluing the
        left column's line to the right column's. Here that would invent
        "Deep Learning" from "hiking, deep" + "Learning & development lead"."""
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        top = pdf.get_y()
        pdf.set_xy(10, top)
        pdf.multi_cell(80, 6, "SKILLS\nPython, SQL, Docker\nInterests: hiking, deep")
        pdf.set_xy(110, top)
        pdf.multi_cell(80, 6, "EXPERIENCE\nBuilt RAG systems on AWS\nLearning & development lead")

        found = set(extract_skills(resume_to_text("resume.pdf", bytes(pdf.output())).text))
        assert {"Python", "SQL", "Docker", "RAG", "AWS"} <= found
        assert "Deep Learning" not in found
